# -*- coding: utf-8 -*-
"""
从 profile.json 生成 Word 简历

用法:
    python build_resume.py [profile.json] [-o 输出.docx]

profile.json 结构见 profile.example.json。
个人数据放在 profile.json（已被 .gitignore 排除），脚本本身可安全公开。
"""
import json
import os
import sys

from docx import Document
from docx.shared import Pt, Cm, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn
from docx.oxml import OxmlElement

FONT = 'Microsoft YaHei'
ACCENT = RGBColor(0x1F, 0x3A, 0x5F)
GREY = RGBColor(0x55, 0x55, 0x55)
RED = RGBColor(0xB0, 0x30, 0x30)

HERE = os.path.dirname(os.path.abspath(__file__))


# ==================== 文档基础设置 ====================

def new_doc():
    doc = Document()
    for s in doc.sections:
        s.top_margin = Cm(1.5)
        s.bottom_margin = Cm(1.4)
        s.left_margin = Cm(1.8)
        s.right_margin = Cm(1.8)
    st = doc.styles['Normal']
    st.font.name = FONT
    st.font.size = Pt(9.5)
    st.element.rPr.rFonts.set(qn('w:eastAsia'), FONT)
    st.paragraph_format.space_after = Pt(2)
    st.paragraph_format.line_spacing = 1.12
    return doc


def rf(run, size=9.5, bold=False, color=None, italic=False):
    """统一设置字体，含中文字体（eastAsia），否则 Word 里中文会回退成宋体"""
    run.font.name = FONT
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.italic = italic
    if color is not None:
        run.font.color.rgb = color
    run._element.rPr.rFonts.set(qn('w:eastAsia'), FONT)
    return run


def para(doc, text='', size=9.5, bold=False, color=None, align=None,
         before=0, after=2, indent=0, italic=False):
    p = doc.add_paragraph()
    if align is not None:
        p.alignment = align
    pf = p.paragraph_format
    pf.space_before = Pt(before)
    pf.space_after = Pt(after)
    if indent:
        pf.left_indent = Cm(indent)
    if text:
        rf(p.add_run(text), size, bold, color, italic)
    return p


def section(doc, title):
    """带下划线的区块标题"""
    p = doc.add_paragraph()
    pf = p.paragraph_format
    pf.space_before = Pt(10)
    pf.space_after = Pt(4)
    rf(p.add_run(title), 12, True, ACCENT)
    pPr = p._element.get_or_add_pPr()
    bdr = OxmlElement('w:pBdr')
    bot = OxmlElement('w:bottom')
    bot.set(qn('w:val'), 'single')
    bot.set(qn('w:sz'), '8')
    bot.set(qn('w:space'), '2')
    bot.set(qn('w:color'), '1F3A5F')
    bdr.append(bot)
    pPr.append(bdr)
    return p


def bullet(doc, text, prefix=None, indent=0.3):
    p = doc.add_paragraph()
    pf = p.paragraph_format
    pf.left_indent = Cm(indent + 0.34)
    pf.first_line_indent = Cm(-0.34)
    pf.space_after = Pt(1.5)
    rf(p.add_run('• '), 9.5, False, ACCENT)
    if prefix:
        rf(p.add_run(prefix), 9.5, True)
    rf(p.add_run(text), 9.5)
    return p


def entry(doc, left, right, left_bold=True, size=10.5, before=7):
    """左标题 + 右对齐时间（用 tab stop 实现）"""
    p = doc.add_paragraph()
    pf = p.paragraph_format
    pf.space_before = Pt(before)
    pf.space_after = Pt(2.5)
    rf(p.add_run(left), size, left_bold)
    p.add_run('\t')
    rf(p.add_run(right), 9, False, GREY)
    pPr = p._element.get_or_add_pPr()
    tabs = OxmlElement('w:tabs')
    tab = OxmlElement('w:tab')
    tab.set(qn('w:val'), 'right')
    tab.set(qn('w:pos'), '9900')
    tabs.append(tab)
    pPr.append(tabs)
    return p


def sub(doc, title, size=10, before=4, indent=0.3):
    return para(doc, title, size, True, None, before=before, after=1.5, indent=indent)


# ==================== 各区块渲染 ====================

def render_header(doc, data):
    para(doc, data['name'], 20, True, ACCENT, WD_ALIGN_PARAGRAPH.CENTER, after=1)
    para(doc, data.get('headline', ''), 10, False, GREY, WD_ALIGN_PARAGRAPH.CENTER, after=1)
    para(doc, data.get('contact', ''), 9, False, GREY, WD_ALIGN_PARAGRAPH.CENTER, after=3)


def render_summary(doc, data):
    s = data.get('summary')
    if not s:
        return
    section(doc, '个人优势')
    if s.get('intro'):
        para(doc, s['intro'], 9.5, after=3)
    for b in s.get('bullets', []):
        bullet(doc, b)


def render_skills(doc, data):
    skills = data.get('skills', [])
    if not skills:
        return
    section(doc, '技术栈')
    for item in skills:
        bullet(doc, item['text'], prefix=item.get('prefix'))


def render_experience(doc, data):
    exps = data.get('experience', [])
    if not exps:
        return
    section(doc, '工作经历')

    for exp in exps:
        entry(doc, f"{exp['company']}　|　{exp['title']}", exp['period'])
        for b in exp.get('bullets', []):
            bullet(doc, b)

        for proj in exp.get('projects', []):
            sub(doc, proj['name'], before=5)
            if proj.get('tech'):
                para(doc, f"技术方案：{proj['tech']}", 8.5, False, GREY, after=2, indent=0.3)
            for b in proj.get('bullets', []):
                bullet(doc, b)
            if proj.get('highlight'):
                para(doc, proj['highlight'], 9, False, RED, after=2, indent=0.3)


def render_education(doc, data):
    edu = data.get('education')
    if not edu:
        return
    section(doc, '教育经历')
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(2)
    rf(p.add_run(edu['school']), 10.5, True)
    rf(p.add_run(f"　|　{edu['line']}"), 9.5, False, GREY)


# ==================== 入口 ====================

def build(profile_path, output_path=None):
    with open(profile_path, encoding='utf-8') as f:
        data = json.load(f)

    doc = new_doc()
    render_header(doc, data)
    render_summary(doc, data)
    render_skills(doc, data)
    render_experience(doc, data)
    render_education(doc, data)

    out = output_path or data.get('output') or os.path.join(HERE, 'resume.docx')
    os.makedirs(os.path.dirname(os.path.abspath(out)), exist_ok=True)
    doc.save(out)
    print(f'OK -> {out}')
    return out


def main():
    args = sys.argv[1:]
    output = None
    if '-o' in args:
        i = args.index('-o')
        output = args[i + 1]
        del args[i:i + 2]

    profile = args[0] if args else os.path.join(HERE, 'profile.json')
    if not os.path.exists(profile):
        print(f'找不到 {profile}')
        print('请复制 profile.example.json 为 profile.json 并填入自己的信息。')
        sys.exit(1)

    build(profile, output)


if __name__ == '__main__':
    main()
